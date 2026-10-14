r"""Genera datos sintéticos de avance de obra y gráficos de control.

El flujo predeterminado crea un libro Excel, vuelve a leerlo y genera:
1. Una curva S de avance planificado y real.
2. Un comparativo del avance por partidas a la fecha de corte.

Uso:
    python generar_curva_s.py
    python generar_curva_s.py --solo-graficos
    python generar_curva_s.py --directorio-salida C:\ruta\de\salida
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import PercentFormatter
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLOR_FONDO = "#09070D"
COLOR_PANEL = "#121018"
COLOR_MORADO = "#9D4EDD"
COLOR_MORADO_CLARO = "#E0AAFF"
COLOR_MORADO_OSCURO = "#5A189A"
COLOR_GRILLA = "#39284A"
COLOR_TEXTO = "#F6F0FA"
COLOR_TEXTO_SECUNDARIO = "#C9B6D4"

PERIODOS = 24
PERIODO_CORTE = 18
FECHA_INICIO = "2026-01-05"


PARTIDAS = [
    ("P01", "Obras preliminares", 180_000, 1, 3, 0, 3),
    ("P02", "Movimiento de tierras", 320_000, 2, 5, 1, 5),
    ("P03", "Cimentaciones", 720_000, 4, 7, 1, 8),
    ("P04", "Estructuras de concreto", 1_450_000, 7, 10, 2, 12),
    ("P05", "Albañilería", 620_000, 12, 7, 1, 8),
    ("P06", "Instalaciones sanitarias", 430_000, 13, 8, 1, 9),
    ("P07", "Instalaciones eléctricas", 480_000, 14, 8, 2, 9),
    ("P08", "Revoques y enlucidos", 360_000, 17, 5, 2, 6),
    ("P09", "Pisos y revestimientos", 510_000, 18, 5, 2, 6),
    ("P10", "Carpintería y vidrios", 390_000, 19, 5, 1, 6),
    ("P11", "Pintura y acabados", 420_000, 20, 5, 2, 6),
    ("P12", "Obras exteriores", 260_000, 21, 4, 1, 5),
]


def _perfil_campana(duracion: int, rng: np.random.Generator | None = None) -> np.ndarray:
    """Distribuye el avance con baja producción al inicio y al cierre."""
    posiciones = np.arange(1, duracion + 1)
    perfil = np.sin(np.pi * posiciones / (duracion + 1)) ** 1.55
    if rng is not None:
        perfil *= rng.uniform(0.92, 1.08, size=duracion)
    return perfil / perfil.sum()


def generar_datos_sinteticos(semilla: int = 42) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Crea cronograma valorizado, avance global y detalle por partida."""
    rng = np.random.default_rng(semilla)
    fechas = pd.date_range(FECHA_INICIO, periods=PERIODOS, freq="14D")
    presupuesto_total = sum(item[2] for item in PARTIDAS)
    detalle: list[dict[str, object]] = []
    resumen_partidas: list[dict[str, object]] = []

    for codigo, nombre, presupuesto, inicio, duracion, retraso, duracion_real in PARTIDAS:
        plan = np.zeros(PERIODOS)
        plan[inicio - 1 : inicio - 1 + duracion] = _perfil_campana(duracion)

        real = np.zeros(PERIODOS)
        inicio_real = inicio - 1 + retraso
        perfil_real = _perfil_campana(duracion_real, rng)
        fin_real = min(PERIODOS, inicio_real + duracion_real)
        real[inicio_real:fin_real] = perfil_real[: fin_real - inicio_real]
        real[PERIODO_CORTE:] = 0.0

        plan_acum = np.cumsum(plan)
        real_acum = np.cumsum(real)
        plan_corte = float(plan_acum[PERIODO_CORTE - 1])
        real_corte = float(real_acum[PERIODO_CORTE - 1])
        variacion = real_corte - plan_corte

        if real_corte >= 0.999:
            estado = "Completada"
        elif real_corte > 0 and variacion < -0.05:
            estado = "En curso - retraso"
        elif real_corte > 0:
            estado = "En curso"
        elif plan_corte > 0:
            estado = "Inicio retrasado"
        else:
            estado = "No iniciada"

        resumen_partidas.append(
            {
                "ID": codigo,
                "Partida": nombre,
                "Presupuesto (S/)": presupuesto,
                "Peso": presupuesto / presupuesto_total,
                "Inicio plan": inicio,
                "Fin plan": inicio + duracion - 1,
                "Avance plan al corte": plan_corte,
                "Avance real al corte": real_corte,
                "Variación (pp)": variacion,
                "Estado": estado,
            }
        )

        for indice, fecha in enumerate(fechas):
            detalle.append(
                {
                    "Periodo": indice + 1,
                    "Fecha": fecha,
                    "ID": codigo,
                    "Partida": nombre,
                    "Presupuesto (S/)": presupuesto,
                    "Peso": presupuesto / presupuesto_total,
                    "Avance plan periodo": float(plan[indice]),
                    "Avance real periodo": float(real[indice]) if indice < PERIODO_CORTE else np.nan,
                    "Avance plan acum. partida": float(plan_acum[indice]),
                    "Avance real acum. partida": float(real_acum[indice]) if indice < PERIODO_CORTE else np.nan,
                }
            )

    detalle_df = pd.DataFrame(detalle)
    partidas_df = pd.DataFrame(resumen_partidas)

    detalle_df["Aporte plan"] = detalle_df["Peso"] * detalle_df["Avance plan periodo"]
    detalle_df["Aporte real"] = detalle_df["Peso"] * detalle_df["Avance real periodo"]
    avance = (
        detalle_df.groupby(["Periodo", "Fecha"], as_index=False)
        .agg(
            **{
                "Avance plan periodo": ("Aporte plan", "sum"),
                "Avance real periodo": ("Aporte real", lambda s: s.sum(min_count=1)),
            }
        )
        .sort_values("Periodo")
    )
    avance["Avance plan acumulado"] = avance["Avance plan periodo"].cumsum()
    avance["Avance real acumulado"] = avance["Avance real periodo"].cumsum()
    avance.loc[avance["Periodo"] > PERIODO_CORTE, "Avance real acumulado"] = np.nan
    avance["Valor planificado (S/)"] = avance["Avance plan acumulado"] * presupuesto_total
    avance["Valor ganado (S/)"] = avance["Avance real acumulado"] * presupuesto_total

    detalle_df = detalle_df.drop(columns=["Aporte plan", "Aporte real"])
    return partidas_df, avance, detalle_df


def _ajustar_excel(ruta_excel: Path, fecha_corte: pd.Timestamp) -> None:
    """Aplica un formato oscuro y morado al libro generado por Python."""
    libro = load_workbook(ruta_excel)
    relleno_fondo = PatternFill("solid", fgColor="09070D")
    relleno_cabecera = PatternFill("solid", fgColor="5A189A")
    relleno_titulo = PatternFill("solid", fgColor="120F18")
    fuente = Font(name="Arial", size=10, color="F6F0FA")
    fuente_cabecera = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    borde = Border(bottom=Side(style="thin", color="4B3562"))

    for hoja in libro.worksheets:
        hoja.sheet_view.showGridLines = False
        hoja.freeze_panes = "A2"
        hoja.sheet_properties.tabColor = "7B2CBF"
        for fila in hoja.iter_rows():
            for celda in fila:
                celda.fill = relleno_fondo
                celda.font = fuente
                celda.alignment = Alignment(vertical="center")
        for celda in hoja[1]:
            celda.fill = relleno_cabecera
            celda.font = fuente_cabecera
            celda.alignment = Alignment(horizontal="center", vertical="center")
            celda.border = borde
        hoja.row_dimensions[1].height = 26
        for columna in range(1, hoja.max_column + 1):
            ancho = max(
                len(str(hoja.cell(fila, columna).value or ""))
                for fila in range(1, min(hoja.max_row, 80) + 1)
            )
            hoja.column_dimensions[get_column_letter(columna)].width = min(max(ancho + 2, 12), 34)

    resumen = libro.create_sheet("Resumen", 0)
    resumen.sheet_view.showGridLines = False
    resumen.sheet_properties.tabColor = "C77DFF"
    resumen["A2"] = "Control de avance de obra"
    resumen["A2"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    resumen["A2"].fill = relleno_titulo
    resumen["A4"] = "Fecha de corte"
    resumen["B4"] = fecha_corte.to_pydatetime()
    resumen["A5"] = "Periodos quincenales"
    resumen["B5"] = PERIODOS
    resumen["A6"] = "Partidas"
    resumen["B6"] = len(PARTIDAS)
    resumen["A8"] = "El libro contiene datos sintéticos reproducibles para practicar control de obra."
    for fila in resumen.iter_rows(min_row=1, max_row=10, min_col=1, max_col=6):
        for celda in fila:
            celda.fill = relleno_fondo
            celda.font = fuente
    resumen["A2"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    resumen.column_dimensions["A"].width = 34
    resumen.column_dimensions["B"].width = 22
    resumen["B4"].number_format = "dd/mm/yyyy"
    libro.save(ruta_excel)


def exportar_excel(
    ruta_excel: Path,
    partidas: pd.DataFrame,
    avance: pd.DataFrame,
    detalle: pd.DataFrame,
) -> None:
    """Guarda los datos sintéticos en un libro Excel estructurado."""
    with pd.ExcelWriter(ruta_excel, engine="openpyxl") as escritor:
        partidas.to_excel(escritor, sheet_name="Partidas", index=False)
        avance.to_excel(escritor, sheet_name="Avance_periodico", index=False)
        detalle.to_excel(escritor, sheet_name="Detalle_partidas", index=False)
    fecha_corte = pd.Timestamp(avance.loc[avance["Periodo"] == PERIODO_CORTE, "Fecha"].iloc[0])
    _ajustar_excel(ruta_excel, fecha_corte)


def _estilizar_eje(ax: plt.Axes) -> None:
    ax.set_facecolor(COLOR_PANEL)
    ax.tick_params(colors=COLOR_TEXTO_SECUNDARIO, labelsize=9)
    for borde in ax.spines.values():
        borde.set_color(COLOR_GRILLA)
    ax.grid(True, color=COLOR_GRILLA, alpha=0.55, linewidth=0.7)
    ax.set_axisbelow(True)


def graficar_curva_s(avance: pd.DataFrame, ruta_salida: Path) -> None:
    """Crea la curva S global de avance físico valorizado."""
    plt.style.use("dark_background")
    fig, ax = plt.subplots(figsize=(13, 7), facecolor=COLOR_FONDO)
    _estilizar_eje(ax)

    ax.plot(
        avance["Fecha"],
        avance["Avance plan acumulado"],
        color=COLOR_MORADO_CLARO,
        linewidth=2.7,
        linestyle="--",
        label="Plan acumulado",
    )
    real = avance.dropna(subset=["Avance real acumulado"])
    ax.plot(
        real["Fecha"],
        real["Avance real acumulado"],
        color=COLOR_MORADO,
        linewidth=3.2,
        marker="o",
        markersize=4.5,
        label="Real acumulado",
    )
    ax.fill_between(
        real["Fecha"],
        real["Avance real acumulado"],
        color=COLOR_MORADO_OSCURO,
        alpha=0.22,
    )

    ultimo = real.iloc[-1]
    plan_corte = float(
        avance.loc[avance["Periodo"] == PERIODO_CORTE, "Avance plan acumulado"].iloc[0]
    )
    real_corte = float(ultimo["Avance real acumulado"])
    ax.axvline(ultimo["Fecha"], color="#F3C4FB", linestyle=":", linewidth=1.5, alpha=0.8)
    ax.annotate(
        f"Corte: {ultimo['Fecha']:%d/%m/%Y}\nPlan {plan_corte:.1%} | Real {real_corte:.1%}\nDesviación {(real_corte-plan_corte):.1%}",
        xy=(ultimo["Fecha"], real_corte),
        xytext=(-145, 34),
        textcoords="offset points",
        color=COLOR_TEXTO,
        fontsize=10,
        bbox={"boxstyle": "round,pad=0.5", "facecolor": "#21142D", "edgecolor": COLOR_MORADO},
        arrowprops={"arrowstyle": "->", "color": COLOR_MORADO_CLARO},
    )

    ax.set_title("Curva S del avance físico valorizado", color=COLOR_TEXTO, fontsize=17, weight="bold", pad=16)
    ax.set_xlabel("Fecha", color=COLOR_TEXTO_SECUNDARIO, labelpad=10)
    ax.set_ylabel("Avance acumulado", color=COLOR_TEXTO_SECUNDARIO, labelpad=10)
    ax.set_ylim(0, 1.05)
    ax.yaxis.set_major_formatter(PercentFormatter(1.0))
    ax.legend(loc="upper left", frameon=False, labelcolor=COLOR_TEXTO)
    fig.autofmt_xdate(rotation=35, ha="right")
    fig.tight_layout()
    fig.savefig(ruta_salida, dpi=180, facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)


def graficar_avance_partidas(partidas: pd.DataFrame, ruta_salida: Path) -> None:
    """Compara el avance planificado y real por partida."""
    plt.style.use("dark_background")
    datos = partidas.copy().iloc[::-1]
    y = np.arange(len(datos))
    alto = 0.36

    fig, ax = plt.subplots(figsize=(13, 8.5), facecolor=COLOR_FONDO)
    _estilizar_eje(ax)
    barras_plan = ax.barh(
        y + alto / 2,
        datos["Avance plan al corte"],
        height=alto,
        color="#C77DFF",
        alpha=0.65,
        label="Plan al corte",
    )
    barras_real = ax.barh(
        y - alto / 2,
        datos["Avance real al corte"],
        height=alto,
        color="#7B2CBF",
        label="Real al corte",
    )
    ax.bar_label(barras_real, labels=[f"{v:.0%}" for v in datos["Avance real al corte"]], padding=4, color=COLOR_TEXTO, fontsize=8)

    ax.set_yticks(y, datos["Partida"])
    ax.set_xlim(0, 1.08)
    ax.xaxis.set_major_formatter(PercentFormatter(1.0))
    ax.set_title("Avance de obra por partidas", color=COLOR_TEXTO, fontsize=17, weight="bold", pad=16)
    ax.set_xlabel("Porcentaje de avance", color=COLOR_TEXTO_SECUNDARIO, labelpad=10)
    ax.grid(axis="y", visible=False)
    ax.legend(loc="lower right", frameon=False, labelcolor=COLOR_TEXTO)
    fig.tight_layout()
    fig.savefig(ruta_salida, dpi=180, facecolor=fig.get_facecolor(), bbox_inches="tight")
    plt.close(fig)


def leer_excel(ruta_excel: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Lee del Excel las tablas que alimentan los gráficos."""

    def leer_hoja_flexible(hoja: str, columna_clave: str) -> pd.DataFrame:
        muestra = pd.read_excel(ruta_excel, sheet_name=hoja, header=None, nrows=8)
        fila_cabecera = next(
            (
                indice
                for indice, fila in muestra.iterrows()
                if columna_clave in fila.astype(str).tolist()
            ),
            None,
        )
        if fila_cabecera is None:
            raise ValueError(f"No se encontró la columna '{columna_clave}' en la hoja '{hoja}'.")
        return pd.read_excel(ruta_excel, sheet_name=hoja, header=fila_cabecera)

    partidas = leer_hoja_flexible("Partidas", "Partida")
    avance = leer_hoja_flexible("Avance_periodico", "Fecha")
    avance["Fecha"] = pd.to_datetime(avance["Fecha"])
    return partidas, avance


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera una curva S y avance por partidas.")
    parser.add_argument(
        "--directorio-salida",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="Carpeta donde se guardan el Excel y los gráficos.",
    )
    parser.add_argument("--semilla", type=int, default=42, help="Semilla para reproducir los datos.")
    parser.add_argument(
        "--solo-graficos",
        action="store_true",
        help="Usa el Excel existente sin regenerarlo.",
    )
    args = parser.parse_args()
    salida = args.directorio_salida.resolve()
    salida.mkdir(parents=True, exist_ok=True)
    ruta_excel = salida / "datos_avance_obra.xlsx"

    if not args.solo_graficos:
        partidas, avance, detalle = generar_datos_sinteticos(args.semilla)
        exportar_excel(ruta_excel, partidas, avance, detalle)

    if not ruta_excel.exists():
        raise FileNotFoundError(f"No existe el archivo requerido: {ruta_excel}")

    partidas_excel, avance_excel = leer_excel(ruta_excel)
    graficar_curva_s(avance_excel, salida / "curva_s.png")
    graficar_avance_partidas(partidas_excel, salida / "avance_por_partidas.png")

    print(f"Excel: {ruta_excel}")
    print(f"Curva S: {salida / 'curva_s.png'}")
    print(f"Avance por partidas: {salida / 'avance_por_partidas.png'}")


if __name__ == "__main__":
    main()
